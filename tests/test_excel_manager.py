import os
import sys
import unittest
import tempfile
import shutil
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from excel_manager import ExcelManager, num_to_words_indian, money, gst_breakdown


class TestExcelManager(unittest.TestCase):

    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.master_file = os.path.join(os.path.dirname(__file__), "..", "Ananya Bill.xlsm")

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_num_to_words_indian(self):
        self.assertEqual(num_to_words_indian(0), "Rupees Zero Only")
        self.assertEqual(num_to_words_indian(1), "Rupees One Only")
        self.assertEqual(num_to_words_indian(15), "Rupees Fifteen Only")
        self.assertEqual(num_to_words_indian(100), "Rupees One Hundred Only")
        self.assertEqual(num_to_words_indian(1000), "Rupees One Thousand Only")
        self.assertEqual(num_to_words_indian(100000), "Rupees One Lakh Only")
        self.assertEqual(num_to_words_indian(10000000), "Rupees One Crore Only")
        self.assertEqual(num_to_words_indian(14868), "Rupees Fourteen Thousand Eight Hundred Sixty Eight Only")
        self.assertEqual(num_to_words_indian(12637.80), "Rupees Twelve Thousand Six Hundred Thirty Seven Paise Eighty Only")

    def test_parse_all_agencies(self):
        mgr = ExcelManager(self.master_file)
        agencies = mgr.parse_all_agencies()
        self.assertEqual(len(agencies), 46)

        # Check structure of first agency
        first = agencies[0]
        self.assertIn('sheet_name', first)
        self.assertIn('agency_name', first)
        self.assertIn('vehicles', first)
        self.assertIn('grand_total', first)
        self.assertGreater(len(first['vehicles']), 0)

    def test_generate_updated_workbook_deduplication(self):
        mgr = ExcelManager(self.master_file)
        out_excel = os.path.join(self.test_dir, "test_out.xlsx")

        # Pass loads with both stripped and unstripped keys to test deduplication
        loads_map = {
            'VINAYAKA BHARATGAS': [{'vehicle_no': 'KA 27 A 9934', 'loads': 25.0, 'rate': 630.0}],
            'ADVAITA': [{'vehicle_no': 'KA 19 D 5917', 'loads': 15.0, 'rate': 630.0}]
        }

        updated = mgr.generate_updated_workbook(
            target_month_year="SEPTEMBER 2026",
            target_date_str="05-09-2026",
            starting_inv_no=2100,
            updated_loads_map=loads_map,
            output_file_path=out_excel
        )

        # Ensure exact agency count: 46 (no duplicate sheets created!)
        self.assertEqual(len(updated), 46)

        wb = openpyxl.load_workbook(out_excel, data_only=True)
        # Check sheet count excluding Sheet1
        agency_sheets = [s for s in wb.sheetnames if s != 'Sheet1']
        self.assertEqual(len(agency_sheets), 46)

        # Check that Vinayaka got the updated load
        vinayaka = next(a for a in updated if 'VINAYAKA' in a['sheet_name'])
        self.assertEqual(vinayaka['vehicles'][0]['loads'], 25.0)

        # Check that Note is written in cell A35
        for s in agency_sheets[:5]:
            ws = wb[s]
            self.assertIn("Note : please complete the payment before 10th of this month", str(ws['A35'].value))

    def test_build_agency_invoice_data(self):
        mgr = ExcelManager(self.master_file)
        data = mgr.build_agency_invoice_data(
            sheet_name="MAHALASA",
            target_month_year="OCTOBER 2026",
            target_date_str="05-10-2026",
            invoice_no=1850,
            vehicles_list=[{'vehicle_no': 'KA 19 D 5917', 'loads': 20.0, 'rate': 550.0}]
        )

        self.assertEqual(data['sheet_name'], "MAHALASA")
        self.assertEqual(data['invoice_no'], 1850)
        self.assertEqual(data['date'], "05-10-2026")
        self.assertIn("OCTOBER 2026", data['month_desc'])
        self.assertEqual(data['total'], 11000.0)
        self.assertEqual(data['sgst'], 990.0)
        self.assertEqual(data['cgst'], 990.0)
        self.assertEqual(data['grand_total'], 12980.0)
        self.assertIn("Twelve Thousand", data['in_words'])

    def test_money_rounds_half_up_like_excel(self):
        self.assertEqual(money(112.545), 112.55)   # float round() gives 112.54
        self.assertEqual(money(2.675), 2.68)       # float round() gives 2.67
        self.assertEqual(money(10), 10.0)

    def test_gst_breakdown(self):
        self.assertEqual(gst_breakdown(1250.50), (1250.5, 112.55, 112.55, 1475.6))
        self.assertEqual(gst_breakdown(0), (0.0, 0.0, 0.0, 0.0))

    def _sheet_names(self, n):
        return ExcelManager(self.master_file).parse_all_agencies()[:n]

    def test_auto_invoice_numbers_skip_manually_set_numbers(self):
        mgr = ExcelManager(self.master_file)
        names = [a['sheet_name'] for a in mgr.parse_all_agencies()]
        out = os.path.join(self.test_dir, "out.xlsx")
        # second agency was given 1001 by hand; automatic numbering starts at 1000
        data = mgr.generate_updated_workbook("OCTOBER 2026", "01-10-2026", 1000, {}, out,
                                             custom_inv_map={names[1]: 1001})
        numbers = [a['invoice_no'] for a in data]
        self.assertEqual(numbers[:3], [1000, 1001, 1002])
        self.assertEqual(len(numbers), len(set(numbers)), "invoice numbers must be unique")

    def test_duplicate_manual_invoice_numbers_are_rejected(self):
        mgr = ExcelManager(self.master_file)
        names = [a['sheet_name'] for a in mgr.parse_all_agencies()]
        out = os.path.join(self.test_dir, "out.xlsx")
        with self.assertRaises(ValueError) as ctx:
            mgr.generate_updated_workbook("OCTOBER 2026", "01-10-2026", 1000, {}, out,
                                          custom_inv_map={names[0]: 5000, names[1]: 5000})
        self.assertIn("5000", str(ctx.exception))
        self.assertFalse(os.path.exists(out), "nothing is saved when numbers clash")

    def _workbook_with_gap(self):
        """Copy of the master where the first agency has vehicles in rows 17 and 19 and row 18 empty."""
        src = openpyxl.load_workbook(self.master_file, keep_vba=False)
        name = [s for s in src.sheetnames if s != 'Sheet1'][0]
        ws = src[name]
        for r in range(17, 30):
            for c in range(1, 7):
                ws.cell(r, c).value = None
        for r, (sl, vno, loads, rate) in {17: (1, "KA01AA1111", 10, 100), 19: (2, "KA01BB2222", 5, 200)}.items():
            ws.cell(r, 1).value, ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value = sl, vno, loads, rate
        path = os.path.join(self.test_dir, "gap.xlsx")
        src.save(path)
        return path, name

    def test_new_vehicles_never_overwrite_existing_ones(self):
        path, name = self._workbook_with_gap()
        vehicles = [{"vehicle_no": "KA01AA1111", "loads": 10, "rate": 100},
                    {"vehicle_no": "KA01BB2222", "loads": 5, "rate": 200},
                    {"vehicle_no": "NEW1", "loads": 1, "rate": 50},
                    {"vehicle_no": "NEW2", "loads": 2, "rate": 50}]
        out = os.path.join(self.test_dir, "out.xlsx")
        data = ExcelManager(path).generate_updated_workbook("OCTOBER 2026", "01-10-2026", 1000, {name: vehicles}, out)
        agency = next(a for a in data if a['sheet_name'] == name)
        self.assertEqual(sorted(v['vehicle_no'] for v in agency['vehicles']),
                         ["KA01AA1111", "KA01BB2222", "NEW1", "NEW2"])
        self.assertEqual(agency['total'], 1000 + 1000 + 50 + 100)
        ws = openpyxl.load_workbook(out)[name]
        self.assertEqual(ws.cell(19, 3).value, "KA01BB2222", "existing vehicle below the gap must survive")
        self.assertEqual({ws.cell(18, 3).value, ws.cell(20, 3).value}, {"NEW1", "NEW2"})

    def test_new_vehicles_fill_every_free_row(self):
        path, name = self._workbook_with_gap()
        vehicles = [{"vehicle_no": "KA01AA1111", "loads": 1, "rate": 1},
                    {"vehicle_no": "KA01BB2222", "loads": 1, "rate": 1}] +                    [{"vehicle_no": f"NEW{i}", "loads": 1, "rate": 1} for i in range(11)]  # 13 total, 11 rows free
        out = os.path.join(self.test_dir, "out.xlsx")
        data = ExcelManager(path).generate_updated_workbook("OCTOBER 2026", "01-10-2026", 1000, {name: vehicles}, out)
        self.assertEqual(len(next(a for a in data if a['sheet_name'] == name)['vehicles']), 13)

    def test_new_vehicles_without_room_raise_instead_of_vanishing(self):
        path, name = self._workbook_with_gap()  # 2 vehicles already on the sheet -> 11 free rows
        vehicles = [{"vehicle_no": f"NEW{i}", "loads": 1, "rate": 1} for i in range(12)]
        out = os.path.join(self.test_dir, "out.xlsx")
        with self.assertRaises(ValueError) as ctx:
            ExcelManager(path).generate_updated_workbook("OCTOBER 2026", "01-10-2026", 1000, {name: vehicles}, out)
        self.assertIn("no room", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
