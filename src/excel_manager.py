import sys
# Force openpyxl to use pure-Python numeric types (int, float, Decimal)
# and completely prevent dummy/incomplete numpy in frozen executables
sys.modules['numpy'] = None

import os
import openpyxl
from openpyxl.styles import Font
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

def num_to_words_indian(num):
    """Converts a number to Indian Rupees format string."""
    if num is None:
        return ""
    try:
        amount = Decimal(str(num)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Invalid currency amount: {num!r}") from exc
    if not amount.is_finite() or amount < 0:
        raise ValueError("Currency amount must be a finite, non-negative number")

    rupees = int(amount)
    paise = int((amount - rupees) * 100)

    units = ['', 'One', 'Two', 'Three', 'Four', 'Five', 'Six', 'Seven', 'Eight', 'Nine', 'Ten', 
             'Eleven', 'Twelve', 'Thirteen', 'Fourteen', 'Fifteen', 'Sixteen', 'Seventeen', 'Eighteen', 'Nineteen']
    tens = ['', '', 'Twenty', 'Thirty', 'Forty', 'Fifty', 'Sixty', 'Seventy', 'Eighty', 'Ninety']

    def convert_below_thousand(n):
        if n == 0: return ''
        if n < 20: return units[n]
        if n < 100:
            return f"{tens[n // 10]} {units[n % 10]}".strip()
        remainder = n % 100
        return f"{units[n // 100]} Hundred {convert_below_thousand(remainder)}".strip()

    def convert_number(n):
        if n == 0: return 'Zero'
        res = []
        if n >= 10000000:
            res.extend([convert_number(n // 10000000), 'Crore'])
            n %= 10000000
        if n >= 100000:
            res.extend([convert_below_thousand(n // 100000), 'Lakh'])
            n %= 100000
        if n >= 1000:
            res.extend([convert_below_thousand(n // 1000), 'Thousand'])
            n %= 1000
        if n > 0:
            res.append(convert_below_thousand(n))
        return ' '.join(res).strip()

    r_text = convert_number(rupees)
    if paise > 0:
        p_text = convert_below_thousand(paise)
        return f"Rupees {r_text} Paise {p_text} Only"
    else:
        return f"Rupees {r_text} Only"


def money(value):
    """Rounds to the paisa half-up, like Excel's ROUND (float round() gives 112.545 -> 112.54)."""
    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def gst_breakdown(subtotal):
    """(subtotal, sgst 9%, cgst 9%, grand total), each rounded to the paisa."""
    subtotal = money(subtotal)
    sgst = money(Decimal(str(subtotal)) * Decimal("0.09"))
    cgst = sgst
    return subtotal, sgst, cgst, money(Decimal(str(subtotal)) + Decimal(str(sgst)) + Decimal(str(cgst)))


VEHICLE_ROWS = range(17, 30)  # 13 vehicle rows in the invoice template


class ExcelManager:
    def __init__(self, file_path):
        self.file_path = file_path
        self.last_output_file_path = None

    def parse_all_agencies(self):
        """Parses all agency invoice sheets from the master workbook."""
        wb = openpyxl.load_workbook(self.file_path, data_only=True)
        agencies = []

        for name in wb.sheetnames:
            if name == 'Sheet1':
                continue
            ws = wb[name]
            
            date_val = ws['F6'].value
            inv_no = ws['F10'].value
            month_desc = ws['A13'].value or ""

            # Agency recipient info
            agency_name = ws['A7'].value or name
            pan_no = ws['F7'].value or ""
            addr_1 = ws['A8'].value or ""
            vendor_code = ws['F8'].value or ""
            addr_2 = ws['A9'].value or ""
            gst_no = ws['F9'].value or ""
            addr_3 = ws['A10'].value or ""
            gstin = ws['A11'].value or ""

            vehicles = []
            for r in range(17, 30):
                sl = ws.cell(r, 1).value
                hsn = ws.cell(r, 2).value or ""
                desc = ws.cell(r, 3).value
                loads = ws.cell(r, 4).value or 0
                rate = ws.cell(r, 5).value or 0
                amt = ws.cell(r, 6).value or 0

                if desc:
                    vehicles.append({
                        'row': r,
                        'sl': sl,
                        'hsn': hsn,
                        'vehicle_no': str(desc).strip(),
                        'loads': loads,
                        'rate': rate,
                        'total_amount': amt
                    })

            total = ws['F30'].value or 0
            sgst = ws['F31'].value or 0
            cgst = ws['F32'].value or 0
            grand_total = ws['F33'].value or 0
            in_words = ws['A34'].value or ""

            agencies.append({
                'sheet_name': name,
                'agency_name': agency_name,
                'pan_no': pan_no,
                'vendor_code': vendor_code,
                'gst_no': gst_no,
                'gstin': gstin,
                'address': [addr_1, addr_2, addr_3],
                'date': date_val,
                'invoice_no': inv_no,
                'month_desc': month_desc,
                'vehicles': vehicles,
                'total': total,
                'sgst': sgst,
                'cgst': cgst,
                'grand_total': grand_total,
                'in_words': in_words
            })

        return agencies

    def generate_updated_workbook(self, target_month_year, target_date_str, starting_inv_no, updated_loads_map, output_file_path, agency_meta_map=None, custom_inv_map=None):
        """
        Generates a new updated Excel file for next month.
        - target_month_year: e.g. "AUGUST 2026"
        - target_date_str: e.g. "01-09-2026"
        - starting_inv_no: int starting counter (e.g. 1721)
        - updated_loads_map: dict of {sheet_name: list_of_vehicle_dicts}
        - agency_meta_map: optional dict containing metadata for new agencies {sheet_name: {agency_name, pan_no, vendor_code, gst_no, gstin, address}}
        - custom_inv_map: optional dict of {sheet_name: custom_inv_no} for manually entered invoice numbers
        """
        os.makedirs(os.path.dirname(os.path.abspath(output_file_path)), exist_ok=True)
        wb = openpyxl.load_workbook(self.file_path, keep_vba=False)
        agency_meta_map = agency_meta_map or {}
        custom_inv_map = custom_inv_map or {}

        # Reference template sheet for copying when new agency added
        ref_template_name = [s for s in wb.sheetnames if s != 'Sheet1'][0]

        # Determine all target sheet names
        # Clean target sheet names: deduplicate matching by stripped lowercase name
        target_sheet_names = []
        seen_sheet_keys = set()
        for sname in wb.sheetnames:
            if sname == 'Sheet1':
                continue
            clean_k = sname.strip().lower()
            if clean_k not in seen_sheet_keys:
                seen_sheet_keys.add(clean_k)
                target_sheet_names.append(sname)

        # Add any newly added agencies from updated_loads_map
        for sname in updated_loads_map.keys():
            if sname == 'Sheet1':
                continue
            clean_k = sname.strip().lower()
            if clean_k not in seen_sheet_keys:
                seen_sheet_keys.add(clean_k)
                target_sheet_names.append(sname)

        current_inv_no = int(starting_inv_no)
        agency_data_list = []
        # Numbers already given to agencies by hand; automatic numbering must not reuse them.
        taken_inv_nos = set()
        for v in custom_inv_map.values():
            try:
                taken_inv_nos.add(int(v))
            except (TypeError, ValueError):
                pass
        assigned_by_sheet = {}

        for name in target_sheet_names:
            if name == 'Sheet1':
                continue

            if name in wb.sheetnames:
                ws = wb[name]
            else:
                # Copy template sheet for newly added agency
                ws = wb.copy_worksheet(wb[ref_template_name])
                ws.title = name
                
                meta = agency_meta_map.get(name, {}) or agency_meta_map.get(name.strip(), {})
                if meta:
                    ws['A7'] = meta.get('agency_name', name)
                    ws['F7'] = meta.get('pan_no', '')
                    addr = meta.get('address', ['', '', ''])
                    ws['A8'] = addr[0] if len(addr) > 0 else ''
                    ws['F8'] = meta.get('vendor_code', '')
                    ws['A9'] = addr[1] if len(addr) > 1 else ''
                    ws['F9'] = meta.get('gst_no', '')
                    ws['A10'] = addr[2] if len(addr) > 2 else ''
                    ws['A11'] = meta.get('gstin', '')

                # Clear vehicle rows 17-29
                for r in range(17, 30):
                    for c in range(1, 7):
                        ws.cell(r, c).value = None

            # 1. Update Date
            ws['F6'] = target_date_str

            # 2. Update Invoice Number (use manually entered number if provided)
            if custom_inv_map and (name in custom_inv_map or name.strip() in custom_inv_map):
                assigned_inv_no = int(custom_inv_map.get(name) or custom_inv_map.get(name.strip()))
            else:
                while current_inv_no in taken_inv_nos:
                    current_inv_no += 1
                assigned_inv_no = current_inv_no
                current_inv_no += 1
            taken_inv_nos.add(assigned_inv_no)
            assigned_by_sheet[name] = assigned_inv_no
            ws['F10'] = assigned_inv_no

            # 3. Update Month Description
            ws['A13'] = f"LPG CYLINDER LOADING AND UNLOADING CHARGES FOR THE MONTH {target_month_year.upper()}"

            # 4. Update Vehicles Loads & Amounts
            # Robust lookup in updated_loads_map (exact, stripped, or case-insensitive)
            agency_vehicles_data = updated_loads_map.get(name)
            if agency_vehicles_data is None:
                agency_vehicles_data = updated_loads_map.get(name.strip())
            if agency_vehicles_data is None:
                name_clean = name.strip().lower()
                for k, v in updated_loads_map.items():
                    if k.strip().lower() == name_clean:
                        agency_vehicles_data = v
                        break
            if agency_vehicles_data is None:
                agency_vehicles_data = []
            
            v_map = {}
            if isinstance(agency_vehicles_data, dict):
                for vno, ld in agency_vehicles_data.items():
                    v_str = str(vno).strip()
                    norm_k = re.sub(r'[\s-]+', '', v_str).upper()
                    item = {'loads': float(ld), 'vehicle_no': v_str}
                    v_map[v_str] = item
                    v_map[norm_k] = item
            else:
                for vitem in agency_vehicles_data:
                    v_str = str(vitem.get('vehicle_no', '')).strip()
                    norm_k = re.sub(r'[\s-]+', '', v_str).upper()
                    v_map[v_str] = vitem
                    v_map[norm_k] = vitem

            subtotal = 0.0
            vehicles_info = []

            # First pass: update existing vehicle rows
            used_vehicles = set()
            max_sl = 0
            empty_rows = []

            for r in VEHICLE_ROWS:
                desc = ws.cell(r, 3).value
                sl_val = ws.cell(r, 1).value
                if sl_val and str(sl_val).isdigit():
                    max_sl = max(max_sl, int(sl_val))

                if desc:
                    v_no = str(desc).strip()
                    norm_v_no = re.sub(r'[\s-]+', '', v_no).upper()
                    used_vehicles.add(v_no)
                    used_vehicles.add(norm_v_no)

                    vdata = v_map.get(v_no) or v_map.get(norm_v_no)
                    if vdata:
                        new_load = float(vdata.get('loads', 0))
                        ws.cell(r, 4).value = new_load

                        if 'rate' in vdata and vdata['rate'] is not None:
                            ws.cell(r, 5).value = float(vdata['rate'])
                    else:
                        new_load = float(ws.cell(r, 4).value or 0)

                    rate = float(ws.cell(r, 5).value or 0)
                    v_total = money(new_load * rate)
                    ws.cell(r, 6).value = v_total
                    subtotal += v_total

                    vehicles_info.append({
                        'sl': ws.cell(r, 1).value,
                        'hsn': ws.cell(r, 2).value or "",
                        'vehicle_no': v_no,
                        'loads': new_load,
                        'rate': rate,
                        'total_amount': v_total
                    })
                else:
                    empty_rows.append(r)

            # Second pass: append any NEW vehicles added by user
            if isinstance(agency_vehicles_data, list):
                if len(agency_vehicles_data) > 13:
                    raise ValueError(
                        f"Agency '{name}' has {len(agency_vehicles_data)} vehicles; "
                        "the workbook template supports a maximum of 13."
                    )
                new_vehicles = []
                for vitem in agency_vehicles_data:
                    v_no = str(vitem['vehicle_no']).strip()
                    norm_v = re.sub(r'[\s-]+', '', v_no).upper()
                    if v_no and v_no not in used_vehicles and norm_v not in used_vehicles:
                        new_vehicles.append(vitem)
                        used_vehicles.update((v_no, norm_v))
                if len(new_vehicles) > len(empty_rows):
                    raise ValueError(
                        f"Agency '{name}' has no room for {len(new_vehicles)} new vehicle(s): only "
                        f"{len(empty_rows)} empty row(s) left in the 13-row invoice template."
                    )
                for vitem, r in zip(new_vehicles, empty_rows):
                    # Rows below a gap may hold existing vehicles, so only truly empty rows are used.
                    v_no = str(vitem['vehicle_no']).strip()
                    max_sl += 1
                    new_load = float(vitem.get('loads', 0))
                    rate = float(vitem.get('rate', 0))
                    v_total = money(new_load * rate)

                    ws.cell(r, 1).value = max_sl
                    ws.cell(r, 3).value = v_no
                    ws.cell(r, 4).value = new_load
                    ws.cell(r, 5).value = rate
                    ws.cell(r, 6).value = v_total
                    subtotal += v_total

                    vehicles_info.append({
                        'sl': max_sl,
                        'hsn': "",
                        'vehicle_no': v_no,
                        'loads': new_load,
                        'rate': rate,
                        'total_amount': v_total
                    })

            # 5. Calculate Subtotal, SGST, CGST, Grand Total
            subtotal, sgst, cgst, grand_total = gst_breakdown(subtotal)
            in_words = num_to_words_indian(grand_total)

            ws['F30'] = subtotal
            ws['F31'] = sgst
            ws['F32'] = cgst
            ws['F33'] = grand_total
            ws['A34'] = in_words
            ws['A35'] = "Note : please complete the payment before 10th of this month"
            ws['A35'].font = Font(name="Calibri", size=10, bold=True, color="C00000")

            agency_data_list.append({
                'sheet_name': name,
                'agency_name': ws['A7'].value or name,
                'pan_no': ws['F7'].value or "",
                'vendor_code': ws['F8'].value or "",
                'gst_no': ws['F9'].value or "",
                'gstin': ws['A11'].value or "",
                'address': [ws['A8'].value or "", ws['A9'].value or "", ws['A10'].value or ""],
                'date': target_date_str,
                'invoice_no': assigned_inv_no,
                'month_desc': ws['A13'].value,
                'vehicles': vehicles_info,
                'total': subtotal,
                'sgst': sgst,
                'cgst': cgst,
                'grand_total': grand_total,
                'in_words': in_words
            })

        by_number = {}
        for sheet, number in assigned_by_sheet.items():
            by_number.setdefault(number, []).append(sheet)
        clashes = {n: sheets for n, sheets in by_number.items() if len(sheets) > 1}
        if clashes:
            details = "; ".join(f"{n}: {', '.join(sheets)}" for n, sheets in sorted(clashes.items()))
            raise ValueError(f"Duplicate invoice numbers — each invoice needs its own number ({details}).")

        # Save workbook handling file lock gracefully
        try:
            wb.save(output_file_path)
            self.last_output_file_path = output_file_path
        except PermissionError:
            alt_path = output_file_path.replace('.xlsx', '_updated.xlsx')
            wb.save(alt_path)
            self.last_output_file_path = alt_path

        return agency_data_list

    def build_agency_invoice_data(self, sheet_name, target_month_year, target_date_str, invoice_no, vehicles_list=None, agency_meta=None):
        """
        Constructs a complete invoice data dictionary for a single agency without processing other sheets.
        Used for on-demand single PDF generation and direct sharing.
        """
        agency_meta = agency_meta or {}
        agency_name = agency_meta.get('agency_name') or sheet_name
        pan_no = agency_meta.get('pan_no', '')
        vendor_code = agency_meta.get('vendor_code', '')
        gst_no = agency_meta.get('gst_no', '')
        gstin = agency_meta.get('gstin', '')
        address = agency_meta.get('address', ['', '', ''])

        # If metadata is incomplete and workbook file exists, look up sheet directly
        if not (pan_no and gst_no) and os.path.exists(self.file_path):
            try:
                wb = openpyxl.load_workbook(self.file_path, data_only=True)
                sname_to_use = sheet_name if sheet_name in wb.sheetnames else sheet_name.strip()
                if sname_to_use in wb.sheetnames:
                    ws = wb[sname_to_use]
                    agency_name = agency_name or ws['A7'].value or sheet_name
                    pan_no = pan_no or ws['F7'].value or ""
                    vendor_code = vendor_code or ws['F8'].value or ""
                    gst_no = gst_no or ws['F9'].value or ""
                    gstin = gstin or ws['A11'].value or ""
                    if not any(address):
                        address = [ws['A8'].value or "", ws['A9'].value or "", ws['A10'].value or ""]
            except Exception:
                pass

        if vehicles_list is None and os.path.exists(self.file_path):
            try:
                wb = openpyxl.load_workbook(self.file_path, data_only=True)
                sname_to_use = sheet_name if sheet_name in wb.sheetnames else sheet_name.strip()
                if sname_to_use in wb.sheetnames:
                    ws = wb[sname_to_use]
                    vehicles_list = []
                    for r in range(17, 30):
                        desc = ws.cell(r, 3).value
                        if desc:
                            vehicles_list.append({
                                'sl': ws.cell(r, 1).value,
                                'hsn': ws.cell(r, 2).value or "",
                                'vehicle_no': str(desc).strip(),
                                'loads': ws.cell(r, 4).value or 0,
                                'rate': ws.cell(r, 5).value or 0,
                            })
            except Exception:
                vehicles_list = []

        vehicles_list = vehicles_list or []
        vehicles_info = []
        subtotal = 0.0

        for idx, v in enumerate(vehicles_list):
            sl = v.get('sl') or (idx + 1)
            hsn = v.get('hsn', '')
            v_no = str(v.get('vehicle_no', '')).strip()
            loads = float(v.get('loads', 0))
            rate = float(v.get('rate', 0))
            v_total = money(loads * rate)
            subtotal += v_total
            vehicles_info.append({
                'sl': sl,
                'hsn': hsn,
                'vehicle_no': v_no,
                'loads': loads,
                'rate': rate,
                'total_amount': v_total
            })

        subtotal, sgst, cgst, grand_total = gst_breakdown(subtotal)
        in_words = num_to_words_indian(grand_total)

        month_desc = f"LPG CYLINDER LOADING AND UNLOADING CHARGES FOR THE MONTH {target_month_year.upper()}"

        return {
            'sheet_name': sheet_name,
            'agency_name': agency_name,
            'pan_no': pan_no,
            'vendor_code': vendor_code,
            'gst_no': gst_no,
            'gstin': gstin,
            'address': address,
            'date': target_date_str,
            'invoice_no': invoice_no,
            'month_desc': month_desc,
            'vehicles': vehicles_info,
            'total': subtotal,
            'sgst': sgst,
            'cgst': cgst,
            'grand_total': grand_total,
            'in_words': in_words
        }

